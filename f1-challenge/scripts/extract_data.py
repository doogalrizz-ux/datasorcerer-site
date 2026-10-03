"""
Extract F1 data from Excel into SQL files for Cloudflare D1.
Usage: python extract_data.py "path/to/F1 Tables.xlsx"
SQL files are written next to this script.
"""

import openpyxl
import re
import os
import sys
from pathlib import Path

EXCEL_PATH = sys.argv[1] if len(sys.argv) > 1 else "F1 Tables.xlsx"
OUT_DIR = Path(__file__).resolve().parent

def clean(val):
    if val is None:
        return None
    s = str(val).strip()
    if s in ("\\N", "NULL", "null", ""):
        return None
    return s

def escape(val):
    if val is None:
        return "NULL"
    return "'" + str(val).replace("'", "''") + "'"

def write_sql(path, table, columns, rows):
    with open(path, "w", encoding="utf-8") as f:
        for row in rows:
            vals = ", ".join(escape(v) for v in row)
            f.write(f"INSERT INTO {table} ({', '.join(columns)}) VALUES ({vals});\n")
    print(f"  Wrote {len(rows)} rows to {path}")

def main():
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)

    # ── Circuits ────────────────────────────────────────────────────────
    print("Reading Circuits...")
    ws = wb["gpm906_f1_Circuits"]
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue  # skip header
        circuit_id, circuit_ref, circuit_name, country, meters, _ = row
        rows.append((
            clean(circuit_id),
            clean(circuit_ref),
            clean(circuit_name),
            clean(country),
            clean(meters),
        ))
    write_sql(
        f"{OUT_DIR}/import_circuits.sql",
        "circuits",
        ["circuit_id", "circuit_ref", "circuit_name", "country", "meters_above_sl"],
        rows
    )

    # ── Drivers ─────────────────────────────────────────────────────────
    print("Reading Drivers...")
    ws = wb["GPM906_f1_Drivers"]
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        driver_id, driver_ref, pref_num, code, firstname, lastname, birthdate, nationality = row
        rows.append((
            clean(driver_id),
            clean(driver_ref),
            clean(pref_num),
            clean(code),
            clean(firstname),
            clean(lastname),
            clean(birthdate),
            clean(nationality),
        ))
    write_sql(
        f"{OUT_DIR}/import_drivers.sql",
        "drivers",
        ["driver_id", "driver_ref", "pref_num", "driver_code", "firstname", "lastname", "birthdate", "nationality"],
        rows
    )

    # ── Races ────────────────────────────────────────────────────────────
    print("Reading Races...")
    ws = wb["gpm906_f1_Races"]
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        race_id, race_year, round_num, circuit_id, race_name, race_date, wiki = row
        rows.append((
            clean(race_id),
            clean(race_year),
            clean(round_num),
            clean(circuit_id),
            clean(race_name),
            clean(race_date),
        ))
    write_sql(
        f"{OUT_DIR}/import_races.sql",
        "races",
        ["race_id", "race_year", "round", "circuit_id", "race_name", "race_date"],
        rows
    )

    # ── Results ──────────────────────────────────────────────────────────
    print("Reading Results...")
    ws = wb["gpm906_f1_Results"]
    rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0:
            continue
        (result_id, race_id, driver_id, constructor_id, car_number, grid,
         finish, finish_text, finish_order, points_awarded, laps_completed,
         race_time_diff, milliseconds, fastest_lap, fast_lap_rank,
         fastest_lap_time, fastest_lap_kph, finish_status) = row
        rows.append((
            clean(result_id),
            clean(race_id),
            clean(driver_id),
            clean(constructor_id),
            clean(car_number),
            clean(grid),
            clean(finish),
            clean(finish_text),
            clean(finish_order),
            clean(points_awarded),
            clean(laps_completed),
            clean(finish_status),
        ))
    write_sql(
        f"{OUT_DIR}/import_results.sql",
        "results",
        ["result_id", "race_id", "driver_id", "constructor_id", "car_number",
         "grid", "finish", "finish_text", "finish_order", "points_awarded",
         "laps_completed", "finish_status"],
        rows
    )

    # ── Schema ───────────────────────────────────────────────────────────
    schema = """
CREATE TABLE IF NOT EXISTS circuits (
  circuit_id      INTEGER PRIMARY KEY,
  circuit_ref     TEXT,
  circuit_name    TEXT,
  country         TEXT,
  meters_above_sl INTEGER
);

CREATE TABLE IF NOT EXISTS drivers (
  driver_id   INTEGER PRIMARY KEY,
  driver_ref  TEXT,
  pref_num    TEXT,
  driver_code TEXT,
  firstname   TEXT,
  lastname    TEXT,
  birthdate   TEXT,
  nationality TEXT
);

CREATE TABLE IF NOT EXISTS races (
  race_id    INTEGER PRIMARY KEY,
  race_year  INTEGER,
  round      INTEGER,
  circuit_id INTEGER,
  race_name  TEXT,
  race_date  TEXT
);

CREATE TABLE IF NOT EXISTS results (
  result_id       INTEGER PRIMARY KEY,
  race_id         INTEGER,
  driver_id       INTEGER,
  constructor_id  INTEGER,
  car_number      TEXT,
  grid            INTEGER,
  finish          TEXT,
  finish_text     TEXT,
  finish_order    INTEGER,
  points_awarded  REAL,
  laps_completed  INTEGER,
  finish_status   INTEGER
);
""".strip()

    with open(f"{OUT_DIR}/schema.sql", "w", encoding="utf-8") as f:
        f.write(schema)
    print(f"  Wrote schema.sql")

    print("\nDone.")

if __name__ == "__main__":
    main()
